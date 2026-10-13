"""Build the Sprint 1 report from implementation and real execution evidence."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
NAME='Abdulrahman Kobeissi'
CAT=json.loads((ROOT/'docs/endpoint-catalog.json').read_text())
SCHEMA=json.loads((ROOT/'evidence/database-schema.json').read_text())
SHOTS=sorted(json.loads((ROOT/'evidence/screenshots/index.json').read_text()),key=lambda x:x['index'])
COL=json.loads((ROOT/'postman/Sprint1.postman_collection.json').read_text(encoding='utf-8'))['item']
RUN=json.loads((ROOT/'evidence/newman-results.json').read_text())
CHECK=json.loads((ROOT/'evidence/integration-checks.json').read_text())
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(7)
    for border in s.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for name in ('TOC 1','TOC 2'):
    if name in doc.styles:
        doc.styles[name].font.size=Pt(9)
        doc.styles[name].paragraph_format.space_after=Pt(3)
        doc.styles[name].paragraph_format.space_before=Pt(0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',28),('Heading 1',18),('Heading 2',14),('Heading 3',12),('Caption',10)]:doc.styles[name].font.size=Pt(size)
doc.styles['Caption'].font.italic=False
head=sec.header.paragraphs[0];head.text='SOFTWARE TOOLS LAB  |  SPRINT 1  |  '+NAME.upper();head.style='Normal'
for r in head.runs:r.font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run(NAME+'  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for r in foot.runs:r.font.size=Pt(8)

def p(text,style=None):return doc.add_paragraph(text,style)
def title(text,level=1):return doc.add_heading(text+' ('+NAME+')',level)
def page(text):doc.add_page_break();title(text)
def table(headers,rows,widths=None,size=9):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
    for row in rows:
        for c,value in zip(t.add_row().cells,row):c.text=str(value)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            if i==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EAF0F5');tcpr.append(fill)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.line_spacing=1.0
                for rr in pp.runs:rr.font.size=Pt(size);rr.bold=(i==0)
    return t
def picture(path,caption,maxheight=7.4):
    path=Path(path);w,h=Image.open(path).size
    width=min(6.85,maxheight*w/h)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(path),width=Inches(width))
    p(NAME+' — '+caption,'Caption')
def code(text):
    para=p(text)
    para.paragraph_format.space_after=Pt(8);para.paragraph_format.line_spacing=1
    for r in para.runs:r.font.name='Consolas';r.font.size=Pt(9)

p('Security Incident Detection and Response Management System','Title')
p('Sprint 1 Backend Implementation Report','Subtitle')
p(NAME)
p('Software Tools Lab\n13 October 2026')
p('This report describes a working backend with nine independently running services, role-based permissions, a relational SQLite database, documented APIs, and reproducible Postman tests. It explains the implementation and links each result to evidence captured from the local system.')
table(['Submission component','Included material'],[
('Source','Nine Flask services, SQL schema, launcher and backup utility'),
('API documentation','60 business operations, OpenAPI reference and Postman collection'),
('Verification','173 HTTP requests, 364 Newman assertions and 18 invariant tests'),
('Evidence','Actual response captures, database records, diagrams and Git history')],[1.55,5.2])
# RESTORABLE_REPORT_NOTE_1
p('Repository: https://github.com/Abkob/Abdulrahman-Kobeissi-435L')

page('Table of Contents')
p('Every section is attributed to Abdulrahman Kobeissi. The evidence appendix is organized by the numbered Postman request, followed by database and GitHub captures.')
toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');doc.add_paragraph()._p.append(toc)
p('Use the Word table of contents links to navigate. Page numbers are updated in the delivered document.')

page('Introduction')
title('Problem and scope',2)
p('Security teams need to connect reported incidents with affected assets, analysts, evidence, known threats, and remediation work. Separate spreadsheets make it difficult to preserve ownership and status history. This backend centralizes those records while exposing a distinct HTTP service for each responsibility.')
p('Sprint 1 covers account and role management, assets, incident investigation, threat intelligence, audit logs, reports, vulnerabilities, notifications, and response playbooks. A product frontend and public cloud deployment are outside this sprint. The read-only database evidence viewer is a development utility.')
title('Objectives and results',2)
p('The implementation provides the 54 required operations and six supporting business operations for logout, role assignment and removal, account status, permissions, and asset history. Four roles are enforced before protected operations. Validation rejects malformed data, unauthorized access and invalid workflow changes. Real HTTP tests exercise all nine processes.')
p('The 173-request collection completed with 364 assertions and no failures. A separate 18-test suite checks security and persistence invariants. Live database checks confirm referential integrity, report reconciliation, empty results and verified backups. These results establish the demonstrated local behavior; they are not load-test or production-readiness claims.')
# RESTORABLE_REPORT_NOTE_2

page('System Architecture')
picture(ROOT/'docs/figures/architecture.png','Nine service processes, HTTP integration and shared persistence.',6.2)
p('Each service has its own Flask application and port. run.py supervises the nine processes. All public paths start with /Abdulrahman-Kobeissi. A shared SQLite database retains relational foreign keys across the sprint’s entities. This is process-level service separation with shared persistence, rather than nine isolated databases.')

page('Service Communication')
table(['Caller','HTTP dependency','Reason'],[
('Protected services','Identity introspection','Validate the bearer token and current user state'),
('Assets, incidents, notifications, playbooks','Identity user lookup','Validate owners, assignees and recipients'),
('Incidents and vulnerabilities','Assets lookup','Validate the affected asset before insertion'),
('Reports','Users, assets, incidents, threats, vulnerabilities and logs','Aggregate actual service responses'),
('Mutation event workers','Audit and notification internal APIs','Deliver durable activity records and alerts')],[1.65,2.15,2.95])
p('The ServiceClient class makes HTTP calls with bounded timeouts. Internal endpoints require a local shared service key, which is generated into the private runtime folder. External requests use bearer sessions. If a required service is unavailable, the caller returns HTTP 503 instead of silently inventing a result.')
p('A business transaction writes both its data and an outbox event. A worker in each service retries pending events through HTTP. Consumers identify events by a unique event_id, making retries idempotent. Pending and failed delivery counts are visible in health responses. This design preserves audit and notification intent when a downstream service is temporarily unavailable.')
p('Example: assigning an incident validates the analyst through Identity, stores the assignment and event in one transaction, then creates an Incident Assignment notification through the Notifications service. Reports obtain source records through HTTP before grouping counts.')
title('Operational boundary',2)
p('The services bind to 127.0.0.1. A production deployment would require production WSGI hosting, TLS, service-specific credentials, secret management, pagination and operational monitoring. No existing production data or schema was changed.')

page('Object Oriented Implementation')
table(['Class','Responsibility','Use of the abstraction'],[
('Settings','Database location, key and service ports','One configuration object per app'),
('Database','Connections and transactions','Commits successful operations and rolls back failures'),
('Repository','Parameterized CRUD and approved filters','Shared persistence behavior without user-supplied SQL'),
('Lifecycle','Allowed ordered status transitions','Reused for incidents, vulnerabilities and response actions'),
('ServiceClient','HTTP communication','Encapsulates timeouts and service errors'),
('EventBus','Durable audit and notification events','Queues and retries idempotent delivery'),
('Context','Service dependencies and actor context','Composes persistence, client and events'),
('CatalogService','Common catalog route and CRUD behavior','Base implementation for two catalogs'),
('ThreatService and VulnerabilityService','Specialized field validation and update rules','Override base validation and lifecycle behavior')],[1.7,2.25,2.8])
p('Encapsulation keeps SQL connection handling inside Database and common record operations inside Repository. Catalog inheritance removes duplicated route plumbing while preserving domain-specific validation. Lifecycle composition applies the same ordered-transition rule to three different workflows. Route functions coordinate these objects and return consistent JSON envelopes.')
p('The source layout separates identity.py, assets.py, incidents.py, catalogs.py, audit_reports.py, notifications.py and playbooks.py. app.py creates the selected service and installs common error handling. The split keeps each service’s authorization and validation close to its business behavior.')

SERVICES=[
('users','Users and Identity Management','Accounts, authentication, role assignment and account status. Usernames and email addresses are unique case-insensitively. Public registration creates active employees; only an authenticated administrator can create privileged accounts.','users and sessions; login_attempts throttles repeated failures. Deleted users are hidden and suspended, while historical foreign keys remain valid.','Registration requires full_name, username, email, password, department and role. Login takes login and password. Responses omit hashes and expose only safe user fields. Profile updates accept full_name, email and department.','A duplicate username returns 409. Invalid credentials return 401. Inactive and suspended accounts cannot authenticate. Role changes, deletion and account-status changes revoke sessions. The last active administrator cannot be removed.','001–035'),
('assets','Asset Management','Maintains inventory, ownership, location, criticality and current status. Administrators manage inventory; auditors read it. Analysts can inspect and search only assets related to their assigned incidents.','assets references its owner in users. asset_history stores before and after JSON and retains deletion history. Referenced assets cannot be deleted.','Creation requires name, asset_type, ip_address, department, owner_id and criticality. operating_system, location and status are optional. Single-asset responses also list related incident and vulnerability IDs.','Criticality is a constrained text value: Low, Medium, High or Critical. Searches combine name, type, IP, department and criticality. SQL values are bound parameters. Invalid IPs and unknown owners are rejected.','036–055; 069–071; 081'),
('incidents','Incident Management','Tracks reporting, assignment, notes and investigation status. Employees see their own reports; analysts see assigned incidents; administrators and auditors have broader read access.','incidents references reporter, optional asset and assigned analyst. investigation_notes references incident and note author.','Creation requires title, description, category, severity and asset_id; asset_id may be null. reporter_id is derived from authentication; an explicitly supplied reporter must match the caller. Notes take content.','The sequence is Open → Under Investigation → Mitigated → Resolved → Closed. Assignment to an active analyst is required before advancing. Notes are visible to authorized investigators and auditors, but hidden from employee incident views. Resolution time is generated when resolved.','056–080'),
('threats','Threat Intelligence','Stores indicators and supports investigation searches. Administrators create, update and delete; analysts and auditors read and search.','threats has a composite unique constraint on threat_type and value. High and Critical creations generate an alert to the acting administrator.','Creation takes threat_type, value, description, severity and source. Updates accept description, severity and source. Responses include the record ID and creation timestamp.','IP values are parsed, domains are normalized to lowercase and trailing-dot-free form, and hashes must be valid MD5, SHA1 or SHA256 hexadecimal strings. Duplicate normalized indicators return 409. Searches support type, value, severity and source. Test indicators use reserved laboratory addresses and example.test domains.','082–099'),
('logs','Security Audit Logging','Records security actions and permission violations. Administrators create and delete log records. Auditors inspect and search; analysts have the read permissions described in the service specification.','logs references an optional user; unauthenticated denials may have no user. The database trigger rejects UPDATE operations. Unique event IDs deduplicate delivery.','Creation requires user_id, action, description, severity and related_service. A record includes a server UTC timestamp. Lists and searches return arrays; no matches return an empty array.','Search filters cover user, date, action, severity and related service. Logs preserve the originally recorded values. The handout explicitly requires administrative deletion; a delete itself emits an audit event, while no modification endpoint exists.','151–163'),
('reports','Security Reports','Generates summary, incident, asset, threat, vulnerability and activity reports for administrators and auditors.','The service does not create report tables. It reads source records over the other services’ HTTP APIs and computes counts and groupings.','The six GET endpoints take no JSON body. Each returns a data object with generated_at, source and its report-specific totals or groups.','Incident reports group severity and status; asset reports group type and criticality; threat reports group severity/type and show recent indicators; vulnerability reports group severity/status; activity reports count actions and login attempts. Summary combines the sources. HTTP 503 surfaces a missing dependency.','164–171'),
('vulnerabilities','Vulnerability Management','Tracks asset vulnerabilities through investigation and remediation. Administrators create and delete; analysts may update; auditors read and search.','vulnerabilities references an affected asset. Resolution time is retained when Closed. Critical or High creation emits a notification.','Creation requires title, description, cve, cvss, severity, asset_id and discovery_source. Updates accept description, severity, status, cvss and resolution_date.','CVE syntax is validated, without claiming a lookup in an external CVE database. CVSS must be numeric from 0 to 10. Status follows Open → Under Investigation → In Progress → Resolved → Closed. A supplied resolution time must be timezone-aware, not future and not before creation.','100–119'),
('notifications','Notification Management','Supports manual and automatic alerts, recipient scoping and read tracking. Administrators create and delete; employees and analysts can read their own alerts; auditors have read-only access.','notifications references its recipient and stores status, created_at, read_date and an optional unique event ID.','Creation requires user_id, title, message, notification_type and severity. Marking read uses PUT with no body. It changes Unread to Read and records the first read timestamp.','Repeated marking as read preserves the original read time. Analysts cannot search another user’s notifications. Searches combine recipient, notification type, severity and read status. A known user without notifications receives a successful empty array.','120–135'),
('playbooks','Security Playbooks and Response Actions','Organizes response procedures by incident category and severity, with ordered actions and execution history.','playbooks references its creator. response_actions references its playbook and assignee. action_history retains every recorded progression. Soft deletion preserves execution history.','Create a playbook with name, description, incident_category and severity. Add actions with description and step_number; assigned_user_id is optional. Status updates require status and may include an assignee.','Steps must be consecutive within a playbook. Pending → In Progress → Completed is enforced, and earlier steps must be completed before later ones begin. Analysts may claim or update their own actions. Completion time is server-generated; completed actions cannot be reassigned.','136–150')]
for service,label,responsibility,db,formats,validation,tests in SERVICES:
    page(label)
    p(responsibility)
    eps=[e for e in CAT if e['service']==service]
    rows=[]
    for e in eps:
        op=e['operation'];body=op.get('requestBody',{}).get('content',{}).get('application/json',{}).get('schema',{})
        required=body.get('required',[])
        fields=', '.join(required) if required else ('Optional update fields' if body else 'Path or query only' if op.get('parameters') else 'None')
        rows.append([e['method'],e['path'].replace('/Abdulrahman-Kobeissi',''),fields])
    table(['Method','Endpoint suffix','Required body fields'],rows,[.65,2.65,3.45],8.5)
    p('All routes use the prefix /Abdulrahman-Kobeissi on port '+str(eps[0]['port'])+'.')
    p(formats);p('Persistence. '+db);p('Rules. '+validation)
    p('Evidence: numbered collection requests '+tests+'. Screenshots appear in the evidence appendix; complete JSON responses and parameters are included in docs/API_REFERENCE.md and docs/openapi.json.')

page('Database Design')
picture(ROOT/'docs/figures/entity-relationships.png','Core entity relationships and foreign-key direction.',6.3)
p('SQLite enforces foreign keys on each connection. IDs are primary keys; session and outbox identifiers use text keys. Timestamps are stored in UTC ISO format. One user may own many assets, report many incidents and receive many notifications. Each note belongs to one incident. An asset may have many incidents and vulnerabilities.')

page('Response and History Relationships')
picture(ROOT/'docs/figures/response-relationships.png','Playbook, action, notification and history relationships.',6.1)
p('Playbooks contain ordered response actions. Users can be assigned many actions. Each action has many immutable historical status entries. Asset history retains a logical asset ID even after deletion, so asset_history.asset_id deliberately does not have a foreign key; actor_id remains linked to users.')
p('User and playbook deletion is soft deletion to preserve relationships and evidence. Asset deletion is restricted if incidents or vulnerabilities reference it. Unique constraints protect usernames, emails, normalized indicator pairs, event IDs and step numbers within a playbook. Enum and range checks provide a second layer beneath API validation.')

purpose={'users':'Accounts and hashed credentials','sessions':'Hashed bearer-token sessions','login_attempts':'Failed-login throttling','assets':'Organizational inventory','asset_history':'Inventory changes before and after mutation','incidents':'Reported security events','investigation_notes':'Analyst evidence and notes','threats':'Threat intelligence indicators','logs':'Audit records','vulnerabilities':'Asset remediation records','notifications':'Recipient alerts and read tracking','playbooks':'Response procedures','response_actions':'Ordered response steps','action_history':'Response execution history','outbox':'Durable HTTP delivery queue'}
groups=[['users','sessions','login_attempts'],['assets','asset_history'],['incidents','investigation_notes'],['threats','logs'],['vulnerabilities','notifications'],['playbooks','response_actions'],['action_history','outbox']]
for group in groups:
    page('Database Dictionary '+str(groups.index(group)+1))
    for name in group:
        title(name.replace('_',' ').title(),2);p(purpose[name]+'.')
        info=SCHEMA[name];fks={f['from']:f['table']+'.'+f['to'] for f in info['foreign_keys']}
        rows=[]
        for c in info['columns']:
            constraint='PK' if c['pk'] else 'NOT NULL' if c['notnull'] else 'Nullable'
            if c['name'] in fks:constraint+='; FK → '+fks[c['name']]
            if c['dflt_value'] is not None:constraint+='; default '+str(c['dflt_value'])
            rows.append([c['name'],c['type'],constraint])
        table(['Field','Type','Key or nullability'],rows,[2.0,.85,3.9],8.5)

page('Error Handling and Validation')
p('Every service uses a common JSON error envelope with status="error", a meaningful message and request_id. Error handlers map expected validation, authentication, authorization, lookup and relationship failures to consistent HTTP statuses. Unhandled failures return a generic server error rather than exposing internals.')
table(['HTTP','Meaning','Demonstrated case'],[(400,'Invalid input','Invalid email, IP, severity, CVSS or unsupported field'),(401,'Authentication failure','Missing token, incorrect login or revoked session'),(403,'Permission or ownership failure','Employee changes incident status; analyst deletes a threat'),(404,'Missing record','Unknown user, asset, incident, catalog entry or playbook'),(409,'Conflict','Duplicate username/indicator or skipped workflow state'),(415,'Unsupported content type','JSON-required operation sent non-JSON content'),(429,'Rate limit','Repeated failed login attempts'),(503,'Dependency unavailable','A report source cannot be contacted')],[.6,2.2,3.95])
p('Input is parsed as JSON objects. Required fields, data types, length bounds, allowed values and unknown keys are checked before persistence. The API rejects empty or malformed requests. Search fields come from a fixed allowlist, and SQL values use placeholders; user input is never inserted into SQL syntax. HTML is not rendered by the backend, so data is stored as text rather than altered with lossy HTML stripping.')
p('Negative requests are intentional test successes when the expected error is returned. Screenshots 013, 015, 025, 051, 065, 072, 085–087, 106, 117 and other appendix cases show these failures. The invariant suite also verifies that SQL-like text remains literal and that a direct log UPDATE fails at the database layer.')

page('API Documentation')
p('docs/openapi.json describes all 60 business operations. docs/API_REFERENCE.md lists purpose, method, complete route, role requirements, path/query parameters, accepted request fields, response examples and possible errors. postman/Sprint1.postman_collection.json contains every numbered test with descriptions, scripts and real response examples.')
p('Protected calls send Authorization: Bearer <token>. Successful reads and updates use HTTP 200; creation uses 201. Data objects and arrays are wrapped in {"status":"success","data":...}. Empty lists remain successful arrays. Errors use the common envelope described in the previous section. Passwords and token values in published examples are redacted.')
title('Registration walkthrough',2)
p('POST /Abdulrahman-Kobeissi/users/register on port 5101 validates the six required fields and uniqueness, hashes the password, then inserts the account and durable events within a transaction. Public callers may register employees; privileged roles require an administrator token. A duplicate is a 409 conflict rather than a second row.')
example=COL[2]['response'][0]
code(example['originalRequest']['body']['raw'])
response=json.loads(example['body']);response['data'].pop('last_login_at',None)
code(json.dumps(response,indent=2))
p('Actual request 003 creates an analyst under administrator authorization. The displayed password is redacted. Figure 003 contains the original Postman view; the collection retains the full safe response.')

page('API Walkthroughs')
title('Incident assignment',2)
p('PUT /Abdulrahman-Kobeissi/incidents/{incident_id}/assign on port 5103 accepts analyst_id. It requires an administrator, checks the incident, calls Identity to verify an active Security Analyst, and stores the assignment. The same transaction queues an audit event and an Incident Assignment notification. Missing incident returns 404, invalid assignee 400, and insufficient permission 403.')
ex=COL[66]['response'][0]
code(ex['originalRequest']['body']['raw'])
data=json.loads(ex['body'])['data']
code(json.dumps({'status':'success','data':{k:data[k] for k in ('id','assigned_analyst_id','status')}},indent=2))
p('The response above is an explicitly shortened extract of request 067; the complete incident fields are in the saved response. Request 065 shows an unauthorized caller rejected from the same operation.')
title('Incident report',2)
p('GET /Abdulrahman-Kobeissi/reports/incidents on port 5106 requires an administrator or auditor and has no body. Reports requests the incident list through HTTP, counts severity and status, and adds generated_at plus the source label. The live reconciliation utility compares those groups with SQLite, not merely with a hardcoded expectation.')
code(COL[164]['response'][0]['body'])
p('This is the saved execution response for request 165. Counts reflect the database at that run. Later reruns create additional uniquely named test records, so current totals may increase without invalidating the dated evidence.')

page('Authentication and Authorization')
p('Werkzeug scrypt hashes passwords with a salt. The database never stores plaintext passwords. Login verifies the hash and issues an opaque random token valid for one hour. Only a SHA-256 digest of that token is stored in sessions. Other services ask Identity to introspect the bearer token and check the current user role and active status.')
table(['Role','Allowed scope','Denied examples'],[
('Security Administrator','Manage users and records, assignments, reports and roles','Cannot bypass data validation or ordered lifecycle rules'),
('Security Analyst','Assigned incidents, related assets, threat reads, vulnerability updates, own notifications and response work','User administration and destructive inventory/catalog operations'),
('Regular Employee/User','Register, maintain own profile, submit and track own incidents, read own notifications','Organizational assets, threats, vulnerabilities, audit logs, reports and playbooks'),
('Auditor','Read users, assets, incidents, catalogs, logs, reports, notifications and playbooks','Any creation, update, deletion or role modification')],[1.45,3.0,2.3])
p('The protected decorator rejects missing/invalid authentication before calling the operation and then checks roles. Resource-specific checks constrain employee ownership and analyst assignments. Denials enqueue an audit record. Removing a privileged role demotes the account to Regular Employee/User. Role and account-status changes revoke existing sessions immediately.')
p('Where the handout’s broad role list and endpoint list differ, the README records the choice: employee and analyst incident lists are scoped to their own/assigned records; analyst asset reads are limited to related incidents; auditor access stays read-only. Login and self-service registration are public entry points, not administrative permissions.')
p('Evidence pairs include login success/failure (001–002), authorized/unauthorized assignment (067/065), deletion pairs across services, inactive/suspended account tests, and token revocation. The test suite checks token storage, session expiry, role changes and protection of the last active administrator.')

page('Verification and Recovery')
table(['Check','Observed result','Evidence file'],[
('Real HTTP scenarios','173 requests; 364 checks; no failures','evidence/http-summary.json'),
('Newman collection execution','173 requests; 364 assertions; no failures','evidence/newman-results.json'),
('Fresh GitHub checkout','173 requests; 364 assertions; no failures on a separate database','evidence/clean-checkout.json'),
('Security and persistence invariants','18 tests pass','evidence/unit-tests.txt'),
('SQLite integrity','ok; zero foreign-key violations','evidence/integration-checks.json'),
('Report reconciliation','All six report types match stored data','evidence/integration-checks.json'),
('Empty responses','Logs and notifications return []','Supplemental-evidence collection'),
('Durable event delivery','Zero pending events at verification','evidence/integration-checks.json'),
('Backup validation','Integrity, foreign keys, counts and SHA-256 checked','Verified local backup manifest')],[1.8,2.25,2.7])
p('The invariant suite uses disposable databases and in-process service clients for focused failure tests. The 173-request HTTP and Newman runs use the actual nine local processes. These are complementary checks; in-process tests are not presented as network evidence.')
p('The supervisor creates a SQLite backup on startup and hourly. The backup utility refuses to overwrite an existing destination, uses the SQLite backup API, validates integrity and foreign keys, and records counts plus a SHA-256 manifest. runtime/backup-status.json makes success or failure visible. Recovery tests open the copied database and verify its data.')
p('Runtime databases, passwords, tokens and backups stay outside Git and the submission archive. The package includes the schema and reproducible sample-data/test procedures. GitHub protects source and sanitized evidence; it does not provide a backup of live application data. These are verified private local recovery copies, not a claimed cloud backup.')

page('GitHub and Version Control')
p('Repository: https://github.com/Abkob/Abdulrahman-Kobeissi-435L')
p('The repository is private under Abkob. The project uses the 435L naming suffix requested by the student. Commits record the implementation, tests, documentation and evidence in meaningful stages.')
history=ROOT/'evidence/git-history.txt'
if history.exists():code(history.read_text(encoding='utf-8').strip())
# RESTORABLE_REPORT_NOTE_4
p('The GitHub evidence appendix shows the actual repository commit page. The final packaging commit may follow that dated screenshot. Reviewers need permission to the private repository.')
title('Reproduction',2)
p('Create a Python virtual environment, install requirements-lock.txt, and run python run.py. The first start creates runtime/bootstrap.json with a random administrator password. Import the collection and clean environment into Postman, select the local environment, set that password locally, and use a local Desktop Agent. Run all requests in order. The README includes exact commands and port assignments.')
p('Reruns use unique names and do not reset an existing database. To start an isolated demonstration, use a separate SECURITY_RUNTIME and SECURITY_DB path. Keep the current database and verified backup before any restoration.')

page('References')
refs=[
('Course specification','Final Project Fall 2026-2027-Sprint 1.docx, supplied course handout.'),
('Prior course work','Local Lab 1, Lab 3, Lab 4 and Lab 5 project materials for Flask, OOP, SQLite, Postman, Git workflow and report conventions.'),
('Flask','Official Quickstart, https://flask.palletsprojects.com/en/stable/quickstart/'),
('Python SQLite','Official sqlite3 documentation, https://docs.python.org/3/library/sqlite3.html'),
('Werkzeug','Official security helper documentation, https://werkzeug.palletsprojects.com/en/stable/utils/'),
('Postman','Official Collection Runner documentation, https://learning.postman.com/docs/tests-and-scripts/running-collections/intro-to-collection-runs/')]
for label,value in refs:p(label+'. '+value)
p('Online documentation was consulted on 13 October 2026. Sample incident descriptions, indicators and account data are local test fixtures, not claims about real attacks or affected people. CVE-formatted sample values demonstrate validation and do not establish real vulnerability findings.')
# RESTORABLE_REPORT_NOTE_3

page('Evidence Appendix')
p('The following screenshots show real response examples captured by Newman from the running local services and opened in Postman. The saved example name identifies the request number and observed status. This is recorded execution evidence, not a mock server or a claim that each screenshot was taken during a fresh Send action.')
p('Crops retain the request title, method, full student-name route, request panel and response/status panel. Long JSON bodies extend beyond the visible editor; the collection and evidence/newman-results.json preserve their complete content. Secrets are redacted. Database captures are fresh read-only queries against runtime/security.db through the included evidence utility.')
p('The request numbers follow the collection’s execution order. Each caption names Abdulrahman Kobeissi and explains the operation. Successful rejection tests are labeled with the actual error status. The table below indexes the groups; individual screenshot captions identify the exact test.')
table(['Request numbers','Evidence group'],[(r,label) for _,label,_,_,_,_,r in SERVICES],[1.6,5.15])
for shot in SHOTS:
    doc.add_page_break()
    name=shot['name'];idx=shot['index'];item=COL[idx-1]
    title('Postman Evidence '+str(idx),2)
    p(name[4:]+' · '+item['request']['method']+' · observed HTTP '+str(item['response'][0]['code']))
    picture(ROOT/'evidence/screenshots'/shot['file'],name+'. The request and saved actual response demonstrate this operation.',7.6)
    if idx in (56,57,58,59):
        p({56:'Severity rationale: Low represents a limited-impact event. This fictional Phishing fixture verifies that the system stores that category correctly.',57:'Severity rationale: Medium represents an event requiring investigation. This fictional Phishing fixture tests that level without asserting a real attack.',58:'Severity rationale: High represents an event affecting important systems or users. This fictional fixture exercises escalation data; no real compromise is claimed.',59:'Severity rationale: Critical represents an event needing immediate response. This fictional fixture tests urgent classification; it is not a measured real-world impact assessment.'}[idx])
    if idx in (89,91,104,111,40,42,124,126):p('Compare this view with the adjacent before/update/after requests. The full response remains available in the collection.')

extra=ROOT/'evidence/screenshots/additional.json'
if extra.exists():
    for shot in json.loads(extra.read_text()):
        doc.add_page_break();title(shot['title'],2)
        picture(ROOT/'evidence/screenshots'/shot['file'],shot['caption'],8.2)
        if shot.get('note'):p(shot['note'])

settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title='Security Incident Detection and Response Management System Sprint 1'
doc.core_properties.author=NAME
doc.core_properties.subject='Nine-service backend implementation and actual execution evidence'
out=ROOT/'Sprint1-Report-Abdulrahman-Kobeissi.docx';doc.save(out)
print(out)
print('Screenshots:',len(SHOTS))
